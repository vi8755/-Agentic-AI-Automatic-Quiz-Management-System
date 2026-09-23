from fastapi import (
    APIRouter,
    Depends,
    HTTPException,
    UploadFile,
    File,
)
from fastapi.responses import StreamingResponse
from io import BytesIO
from openpyxl import Workbook
from ..services.descriptive_evaluation_service import (
    extract_pdf_text,
    extract_questions_from_pdf,
)
from sqlalchemy.orm import Session
from sqlalchemy import exists
from ..security import (
    get_current_user,
    require_role,
    require_teacher,
)
import os
import uuid
from ..tools.email_tools import send_quiz_email
from ..database import SessionLocal

from ..models import (
    User,
    UserRole,
    Teacher,
    Section,
    Student,
    TeacherSection,

    # Descriptive Assignment Models
    DescriptiveAssignment,
    DescriptiveAssignmentQuestion,
    DescriptiveAssignmentSection,
    DescriptiveSubmission,
)

from ..schemas import (
    TeacherCreate,
    TeacherRegistrationCreate,
    TeacherResponse,
    TeacherAssignmentsResponse,
    TeacherMySectionResponse,
    TeacherDashboardResponse,
    TeacherQuizResponse,
    TeacherQuizAssignmentResponse,
    UpdateQuizRequest,
    QuizWithQuestionsResponse,
    UpdateQuestionRequest,
    GenerateQuestionRequest,
    AssignQuizSectionRequest,
    TeacherQuizReportList,
    TeacherStudentReport,
    TeacherProfileResponse,
    TeacherProfileUpdate,
    ChangePasswordRequest,

    # Descriptive Assignment Schemas
    DescriptiveSubmissionEvaluationUpdate,
    DescriptiveAssignmentCreate,
    DescriptiveAssignmentSectionAssign,
)
import secrets
from ..services.teacher_service import (
    create_teacher,
    register_teacher,
    get_teacher_assignments,
    get_my_sections,
    get_teacher_dashboard,
    get_teacher_quizzes,
    get_quiz_assignments,
    get_teacher_analytics,
    get_quiz_performance,
    get_section_performance,
    get_teacher_quiz_by_id,
    get_recent_quiz_activity,
    update_teacher_quiz,
    delete_teacher_quiz,
    duplicate_teacher_quiz,
    publish_teacher_quiz,
    move_quiz_to_draft,
    regenerate_teacher_question,
    update_teacher_question,
    generate_teacher_question,
    get_teacher_profile,
    update_teacher_profile,
    change_teacher_password,
    export_quiz_assignments,

    # Descriptive Assignment Services
    get_teacher_descriptive_assignment_submissions,
    get_teacher_descriptive_submission,
    review_teacher_descriptive_submission,
    get_teacher_quiz_performance,
    get_student_performance,
)

from app.services import teacher_service


# =========================================================
# ROUTER
# =========================================================

router = APIRouter(
    prefix="/teachers",
    tags=["Teachers"],
)


# =========================================================
# DATABASE
# =========================================================

def get_db():
    db = SessionLocal()

    try:
        yield db

    finally:
        db.close()


# =========================================================
# TEACHER ASSIGNMENTS
# =========================================================

@router.get(
    "/{teacher_id}/assignments",
    response_model=TeacherAssignmentsResponse,
)
def get_teacher_assignments_api(
    teacher_id: int,
    db: Session = Depends(get_db),
):
    try:
        return get_teacher_assignments(
            db,
            teacher_id,
        )

    except ValueError as e:
        raise HTTPException(
            status_code=404,
            detail=str(e),
        )


# =========================================================
# MY SECTIONS
# =========================================================

@router.get(
    "/my-sections",
    response_model=list[TeacherMySectionResponse],
)
def my_sections(
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
    db: Session = Depends(get_db),
):
    try:
        return get_my_sections(
            db,
            current_user,
        )

    except ValueError as e:
        raise HTTPException(
            status_code=404,
            detail=str(e),
        )


# =========================================================
# TEACHER DASHBOARD
# =========================================================

@router.get(
    "/dashboard",
    response_model=TeacherDashboardResponse,
)
def teacher_dashboard(
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
    db: Session = Depends(get_db),
):
    try:
        return get_teacher_dashboard(
            db,
            current_user,
        )

    except ValueError as e:
        raise HTTPException(
            status_code=404,
            detail=str(e),
        )


# =========================================================
# REGISTER TEACHER
# =========================================================

@router.post(
    "/register",
    response_model=TeacherResponse,
)
def register_teacher_api(
    teacher: TeacherRegistrationCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_role(UserRole.DEAN)
    ),
):
    try:
        return register_teacher(
            db,
            teacher,
        )

    except ValueError as e:
        raise HTTPException(
            status_code=400,
            detail=str(e),
        )


# =========================================================
# QUIZ MANAGEMENT
# =========================================================

@router.get(
    "/quizzes",
    response_model=list[TeacherQuizResponse],
)
def teacher_quizzes(
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
    db: Session = Depends(get_db),
):
    try:
        return get_teacher_quizzes(
            db,
            current_user,
        )

    except ValueError as e:
        raise HTTPException(
            status_code=404,
            detail=str(e),
        )


# =========================================================
# QUIZ ASSIGNMENTS
# =========================================================

@router.get(
    "/quizzes/{quiz_id}/assignments",
    response_model=list[TeacherQuizAssignmentResponse],
)
def teacher_quiz_assignments(
    quiz_id: int,
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
    db: Session = Depends(get_db),
):
    try:
        return get_quiz_assignments(
            db,
            current_user,
            quiz_id,
        )

    except ValueError as e:
        raise HTTPException(
            status_code=404,
            detail=str(e),
        )

# =========================================================
# EXPORT QUIZ ASSIGNMENTS TO EXCEL
# =========================================================

@router.get(
    "/quizzes/{quiz_id}/assignments/export"
)
def export_quiz_assignments_api(
    quiz_id: int,
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
    db: Session = Depends(get_db),
):

    try:

        data = export_quiz_assignments(
            db,
            current_user,
            quiz_id,
        )

        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = "Quiz Assignments"

        # -----------------------------------------
        # Headers
        # -----------------------------------------

        headers = [
            "Student Name",
            "Roll No",
            "Email",
            "Assigned Date",
            "Due Date",
            "Status",
            "Score",
            "Total Marks",
            "Percentage",
        ]

        worksheet.append(headers)

        # -----------------------------------------
        # Data
        # -----------------------------------------

        for row in data["rows"]:

            worksheet.append([
                row["Student Name"],
                row["Roll No"],
                row["Email"],
                row["Assigned Date"],
                row["Due Date"],
                row["Status"],
                row["Score"],
                row["Total Marks"],
                row["Percentage"],
            ])

        # -----------------------------------------
        # Column widths
        # -----------------------------------------

        worksheet.column_dimensions["A"].width = 25
        worksheet.column_dimensions["B"].width = 15
        worksheet.column_dimensions["C"].width = 30
        worksheet.column_dimensions["D"].width = 20
        worksheet.column_dimensions["E"].width = 20
        worksheet.column_dimensions["F"].width = 15
        worksheet.column_dimensions["G"].width = 12
        worksheet.column_dimensions["H"].width = 15
        worksheet.column_dimensions["I"].width = 15

        # -----------------------------------------
        # Freeze header row
        # -----------------------------------------

        worksheet.freeze_panes = "A2"

        # -----------------------------------------
        # Create Excel file
        # -----------------------------------------

        file_stream = BytesIO()

        workbook.save(file_stream)

        file_stream.seek(0)

        # -----------------------------------------
        # Safe filename
        # -----------------------------------------

        filename = (
            data["quiz_title"]
            .replace(" ", "_")
            .replace("/", "_")
            + "_Assignment_Report.xlsx"
        )

        return StreamingResponse(
            file_stream,
            media_type=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition":
                    f'attachment; filename="{filename}"'
            },
        )

    except ValueError as e:

        raise HTTPException(
            status_code=404,
            detail=str(e),
        )
# =========================================================
# ASSIGN MCQ QUIZ TO SECTION
# =========================================================

@router.post(
    "/quizzes/{quiz_id}/assign",
)
def assign_quiz(
    quiz_id: int,
    request: AssignQuizSectionRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
):
    try:
        return teacher_service.assign_quiz_to_section(
            db=db,
            current_user=current_user,
            quiz_id=quiz_id,
            request=request,
        )

    except ValueError as e:
        raise HTTPException(
            status_code=400,
            detail=str(e),
        )


# =========================================================
# QUIZ DETAILS
# =========================================================

@router.get(
    "/quizzes/{quiz_id}",
    response_model=QuizWithQuestionsResponse,
)
def teacher_quiz_details(
    quiz_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
):
    try:
        return get_teacher_quiz_by_id(
            db,
            current_user,
            quiz_id,
        )

    except ValueError as e:
        raise HTTPException(
            status_code=404,
            detail=str(e),
        )


# =========================================================
# ANALYTICS
# =========================================================

@router.get("/analytics")
def teacher_analytics(
      section_id: int | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
):
    teacher = (
        db.query(Teacher)
        .filter(
            Teacher.user_id == current_user.id
        )
        .first()
    )

    if not teacher:
        raise HTTPException(
            status_code=404,
            detail="Teacher profile not found.",
        )

    return get_teacher_analytics(
        db,
        teacher.id,
        section_id,
    )


@router.get(
    "/analytics/quiz-performance"
)
def teacher_quiz_performance(
    section_id: int | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
):
    teacher = (
        db.query(Teacher)
        .filter(
            Teacher.user_id == current_user.id
        )
        .first()
    )

    if not teacher:
        raise HTTPException(
            status_code=404,
            detail="Teacher profile not found.",
        )

    return get_quiz_performance(
        db,
        teacher.id,
        section_id,
    )


@router.get(
    "/analytics/section-performance"
)
def teacher_section_performance(
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
):
    teacher = (
        db.query(Teacher)
        .filter(
            Teacher.user_id == current_user.id
        )
        .first()
    )

    if not teacher:
        raise HTTPException(
            status_code=404,
            detail="Teacher profile not found.",
        )

    return get_section_performance(
        db,
        teacher.id,
    )


@router.get(
    "/analytics/recent-quizzes"
)
def teacher_recent_quizzes(
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
):
    teacher = (
        db.query(Teacher)
        .filter(
            Teacher.user_id == current_user.id
        )
        .first()
    )

    if not teacher:
        raise HTTPException(
            status_code=404,
            detail="Teacher not found",
        )

    return get_recent_quiz_activity(
        db,
        teacher.id,
    )


# =========================================================
# UPDATE QUIZ
# =========================================================

@router.put(
    "/quizzes/{quiz_id}"
)
def edit_teacher_quiz(
    quiz_id: int,
    quiz_data: UpdateQuizRequest,
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
    db: Session = Depends(get_db),
):
    try:
        return update_teacher_quiz(
            db,
            current_user,
            quiz_id,
            quiz_data,
        )

    except ValueError as e:
        raise HTTPException(
            status_code=404,
            detail=str(e),
        )


# =========================================================
# DELETE QUIZ
# =========================================================

@router.delete(
    "/quizzes/{quiz_id}"
)
def delete_quiz(
    quiz_id: int,
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
    db: Session = Depends(get_db),
):
    try:
        return delete_teacher_quiz(
            db,
            current_user,
            quiz_id,
        )

    except ValueError as e:
        raise HTTPException(
            status_code=404,
            detail=str(e),
        )

    except PermissionError as e:
        raise HTTPException(
            status_code=403,
            detail=str(e),
        )


# =========================================================
# DUPLICATE QUIZ
# =========================================================

@router.post(
    "/quizzes/{quiz_id}/duplicate"
)
def duplicate_quiz(
    quiz_id: int,
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
    db: Session = Depends(get_db),
):
    try:
        return duplicate_teacher_quiz(
            db,
            current_user,
            quiz_id,
        )

    except ValueError as e:
        raise HTTPException(
            status_code=404,
            detail=str(e),
        )

    except PermissionError as e:
        raise HTTPException(
            status_code=403,
            detail=str(e),
        )


# =========================================================
# PUBLISH QUIZ
# =========================================================

@router.patch(
    "/quizzes/{quiz_id}/publish"
)
def publish_quiz(
    quiz_id: int,
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
    db: Session = Depends(get_db),
):
    try:
        return publish_teacher_quiz(
            db,
            current_user,
            quiz_id,
        )

    except ValueError as e:
        raise HTTPException(
            status_code=404,
            detail=str(e),
        )

    except PermissionError as e:
        raise HTTPException(
            status_code=403,
            detail=str(e),
        )


# =========================================================
# MOVE QUIZ TO DRAFT
# =========================================================

@router.patch(
    "/quizzes/{quiz_id}/draft"
)
def draft_quiz(
    quiz_id: int,
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
    db: Session = Depends(get_db),
):
    try:
        return move_quiz_to_draft(
            db,
            current_user,
            quiz_id,
        )

    except ValueError as e:
        raise HTTPException(
            status_code=404,
            detail=str(e),
        )

    except PermissionError as e:
        raise HTTPException(
            status_code=403,
            detail=str(e),
        )


# =========================================================
# QUESTION MANAGEMENT
# =========================================================

@router.post(
    "/quizzes/{quiz_id}/questions/{question_id}/regenerate"
)
def regenerate_question_api(
    quiz_id: int,
    question_id: int,
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
    db: Session = Depends(get_db),
):
    try:
        return regenerate_teacher_question(
            db,
            current_user,
            quiz_id,
            question_id,
        )

    except ValueError as e:
        raise HTTPException(
            status_code=404,
            detail=str(e),
        )


@router.put(
    "/questions/{question_id}"
)
def update_question(
    question_id: int,
    question_data: UpdateQuestionRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
):
    return update_teacher_question(
        db=db,
        current_user=current_user,
        question_id=question_id,
        question_data=question_data,
    )


@router.post(
    "/quizzes/{quiz_id}/questions/generate"
)
def generate_question(
    quiz_id: int,
    request: GenerateQuestionRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
):
    return generate_teacher_question(
        db=db,
        current_user=current_user,
        quiz_id=quiz_id,
        request=request,
    )


# =========================================================
# QUIZ REPORTS
# =========================================================

@router.get(
    "/quizzes/{quiz_id}/reports",
    response_model=TeacherQuizReportList,
)
def teacher_quiz_reports(
    quiz_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
):
    return teacher_service.get_teacher_quiz_reports(
        db,
        current_user,
        quiz_id,
    )


@router.get(
    "/reports/{response_id}",
    response_model=TeacherStudentReport,
)
def teacher_student_report(
    response_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
):
    return teacher_service.get_student_report(
        db,
        current_user,
        response_id,
    )


# =========================================================
# TEACHER PROFILE
# =========================================================

@router.get(
    "/profile",
    response_model=TeacherProfileResponse,
)
def teacher_profile(
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_teacher
    ),
):
    return get_teacher_profile(
        db,
        current_user,
    )


@router.put(
    "/profile"
)
def update_profile(
    profile: TeacherProfileUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_teacher
    ),
):
    return update_teacher_profile(
        db,
        current_user,
        profile,
    )

@router.get(
    "/quizzes/{quiz_id}/performance"
)
def teacher_quiz_performance_details(
    quiz_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
):
    try:

        return get_teacher_quiz_performance(
            db,
            current_user,
            quiz_id,
        )

    except ValueError as e:

        raise HTTPException(
            status_code=404,
            detail=str(e),
        )
@router.get(
    "/quizzes/{quiz_id}/performance/export"
)
def export_teacher_quiz_performance(
    quiz_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
):
    # =========================================================
    # GET PERFORMANCE DATA
    # =========================================================

    try:

        performance = get_teacher_quiz_performance(
            db=db,
            current_user=current_user,
            quiz_id=quiz_id,
        )

    except ValueError as e:

        raise HTTPException(
            status_code=404,
            detail=str(e),
        )

    # =========================================================
    # CREATE WORKBOOK
    # =========================================================

    workbook = Workbook()

    students_sheet = workbook.active

    students_sheet.title = "All Students"

    # =========================================================
    # HEADER
    # =========================================================

    students_sheet.append(
        [
            "Rank",
            "Student Name",
            "Roll No",
            "Email",
            "Section",
            "Score",
            "Total Marks",
            "Percentage",
            "Status",
        ]
    )

    # =========================================================
    # ADD ALL STUDENTS
    # =========================================================

    for student in performance["students"]:

        students_sheet.append(
            [
                student.get(
                    "rank",
                    "",
                ),
                student.get(
                    "student_name",
                    "",
                ),
                student.get(
                    "roll_no",
                    "",
                ),
                student.get(
                    "email",
                    "",
                ),
                student.get(
                    "section",
                    "",
                ),
                student.get(
                    "score",
                    0,
                ),
                student.get(
                    "total_marks",
                    0,
                ),
                student.get(
                    "percentage",
                    0,
                ),
                student.get(
                    "status",
                    "",
                ),
            ]
        )

    # =========================================================
    # FREEZE HEADER ROW
    # =========================================================

    students_sheet.freeze_panes = "A2"

    # =========================================================
    # ENABLE EXCEL FILTER
    # =========================================================

    students_sheet.auto_filter.ref = (
        students_sheet.dimensions
    )

    # =========================================================
    # AUTO COLUMN WIDTH
    # =========================================================

    for column in students_sheet.columns:

        max_length = 0

        column_letter = (
            column[0].column_letter
        )

        for cell in column:

            if cell.value is None:
                continue

            value = str(
                cell.value
            )

            max_length = max(
                max_length,
                len(value),
            )

        students_sheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 3,
            40,
        )

    # =========================================================
    # CREATE EXCEL FILE IN MEMORY
    # =========================================================

    excel_file = BytesIO()

    workbook.save(
        excel_file
    )

    excel_file.seek(0)

    # =========================================================
    # SAFE FILE NAME
    # =========================================================

    quiz_title = str(
        performance["quiz"].get(
            "title",
            f"quiz_{quiz_id}",
        )
    )

    safe_title = "".join(
        character
        if character.isalnum()
        or character in (
            " ",
            "-",
            "_",
        )
        else "_"
        for character in quiz_title
    ).strip()

    if not safe_title:

        safe_title = (
            f"quiz_{quiz_id}"
        )

    filename = (
        f"{safe_title}_performance.xlsx"
    )

    # =========================================================
    # RETURN EXCEL FILE
    # =========================================================

    return StreamingResponse(
        excel_file,
        media_type=(
            "application/vnd."
            "openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )
@router.put(
    "/change-password"
)
def update_password(
    password_data: ChangePasswordRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_teacher
    ),
):
    return change_teacher_password(
        db,
        current_user,
        password_data,
    )


# =========================================================
# DESCRIPTIVE ASSIGNMENT MANAGEMENT
# =========================================================


# =========================================================
# GET TEACHER'S SECTIONS
#
# Used by Create / Assign Descriptive Assignment UI
# =========================================================

@router.get(
    "/descriptive-assignments/sections"
)
def get_teacher_descriptive_sections(
    db: Session = Depends(get_db),
    current_user: User = Depends(
        get_current_user
    ),
):
    teacher = (
        db.query(Teacher)
        .filter(
            Teacher.user_id == current_user.id
        )
        .first()
    )

    if not teacher:
        raise HTTPException(
            status_code=404,
            detail="Teacher not found.",
        )

    sections = (
        db.query(Section)
        .join(
            TeacherSection,
            TeacherSection.section_id == Section.id,
        )
        .filter(
            TeacherSection.teacher_id == teacher.id
        )
        .all()
    )

    return [
        {
            "id": section.id,

            "name": getattr(
                section,
                "name",
                f"Section {section.id}",
            ),

            "section_name": getattr(
                section,
                "section_name",
                None,
            ),

            "department": getattr(
                section,
                "department",
                None,
            ),
        }

        for section in sections
    ]


# =========================================================
# CREATE DESCRIPTIVE ASSIGNMENT
#
# Creates assignment as Draft.
#
# Sections are NOT assigned here.
# Teacher assigns sections using /assign endpoint.
# =========================================================

# =========================================================
# CREATE DESCRIPTIVE ASSIGNMENT
# =========================================================

@router.post(
    "/descriptive-assignments"
)
def create_descriptive_assignment(
    data: DescriptiveAssignmentCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        get_current_user
    ),
):
    # =====================================================
    # FIND TEACHER
    # =====================================================

    teacher = (
        db.query(Teacher)
        .filter(
            Teacher.user_id == current_user.id
        )
        .first()
    )

    if not teacher:
        raise HTTPException(
            status_code=404,
            detail="Teacher not found.",
        )

    # =====================================================
    # VALIDATE ASSIGNMENT TYPE
    # =====================================================

    assignment_type = (
        data.assignment_type.upper()
        if data.assignment_type
        else ""
    )

    if assignment_type not in [
        "MANUAL",
        "PDF",
    ]:
        raise HTTPException(
            status_code=400,
            detail=(
                "Invalid assignment type. "
                "Use MANUAL or PDF."
            ),
        )

     # =====================================================
# MANUAL VALIDATION
# =====================================================

    if assignment_type == "MANUAL":

     if not data.questions:
        raise HTTPException(
            status_code=400,
            detail=(
                "Manual assignment must "
                "contain questions."
            ),
        )

    if not data.duration_minutes:
        raise HTTPException(
            status_code=400,
            detail=(
                "Duration is required "
                "for manual assignments."
            ),
        )

    if data.duration_minutes <= 0:
        raise HTTPException(
            status_code=400,
            detail=(
                "Duration must be greater than 0 minutes."
            ),
        )

    # =====================================================
    # PDF VALIDATION
    # =====================================================

    if assignment_type == "PDF":

        if not data.question_pdf_url:

            raise HTTPException(
                status_code=400,
                detail=(
                    "Question PDF is required "
                    "for PDF assignment."
                ),
            )

    # =====================================================
    # CREATE ASSIGNMENT
    # =====================================================

    assignment = DescriptiveAssignment(
        teacher_id=teacher.id,
        subject_id=data.subject_id,
        title=data.title,
        instructions=data.instructions,
        due_date=data.due_date,
        status="Draft",

        assignment_type=assignment_type,
        duration_minutes=data.duration_minutes,

        question_pdf_url=(
            data.question_pdf_url
            if assignment_type == "PDF"
            else None
        ),

        question_pdf_name=(
            data.question_pdf_name
            if assignment_type == "PDF"
            else None
        ),
    )

    db.add(assignment)

    # Get assignment.id
    db.flush()

    # =====================================================
    # MANUAL QUESTIONS
    # =====================================================

    if assignment_type == "MANUAL":

        for question in data.questions:

            db.add(
                DescriptiveAssignmentQuestion(
                    assignment_id=assignment.id,

                    question_text=(
                        question.question_text
                    ),

                    max_marks=(
                        question.max_marks
                    ),

                    expected_answer=(
                        question.expected_answer
                    ),

                    evaluation_rubric=(
                        question.evaluation_rubric
                    ),

                    question_order=(
                        question.question_order
                    ),
                )
            )

    # =====================================================
    # PDF QUESTIONS
    # =====================================================

    elif assignment_type == "PDF":

        try:

            # -------------------------------------------------
            # IMPORTANT:
            #
            # First extract the ACTUAL TEXT from the PDF.
            #
            # Do NOT pass question_pdf_url directly to
            # extract_questions_from_pdf().
            # -------------------------------------------------

            from ..services.descriptive_evaluation_service import (
                extract_pdf_text,
                extract_questions_from_pdf,
            )

            question_pdf_text = (
                extract_pdf_text(
                    data.question_pdf_url
                )
            )

            if not question_pdf_text:

                raise ValueError(
                    "No readable text could be extracted "
                    "from the question PDF."
                )

            # -------------------------------------------------
            # Now extract actual questions from the text
            # -------------------------------------------------

            pdf_questions = (
                extract_questions_from_pdf(
                    question_pdf_text
                )
            )

            if not pdf_questions:

                raise ValueError(
                    "No questions could be extracted "
                    "from the question PDF."
                )

        except Exception as e:

            db.rollback()

            raise HTTPException(
                status_code=400,
                detail=(
                    "Unable to extract questions from "
                    f"question PDF: {str(e)}"
                ),
            )

        # -------------------------------------------------
        # CREATE DATABASE QUESTION RECORDS
        # -------------------------------------------------

        for question in pdf_questions:

            db.add(
                DescriptiveAssignmentQuestion(
                    assignment_id=assignment.id,

                    question_text=(
                        question[
                            "question_text"
                        ]
                    ),

                    # IMPORTANT:
                    # Use marks extracted from PDF
                    max_marks=(
                        question.get(
                            "max_marks",
                            10,
                        )
                    ),

                    expected_answer=(
                        question.get(
                            "expected_answer"
                        )
                    ),

                    evaluation_rubric=(
                        question.get(
                            "evaluation_rubric"
                        )
                    ),

                    question_order=(
                        question[
                            "question_order"
                        ]
                    ),
                )
            )

    # =====================================================
    # SAVE
    # =====================================================

    db.commit()

    db.refresh(assignment)

    # =====================================================
    # RESPONSE
    # =====================================================

    return {
        "id": assignment.id,

        "title": assignment.title,

        "status": assignment.status,

        "assignment_type": (
            assignment.assignment_type
        ),

        "question_pdf_url": (
            assignment.question_pdf_url
        ),

        "question_pdf_name": (
            assignment.question_pdf_name
        ),

        "message": (
            "Assignment created successfully."
        ),
    }

@router.get(
    "/descriptive-assignments/{assignment_id}"
)
def get_teacher_descriptive_assignment(
    assignment_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        get_current_user
    ),
):
    teacher = (
        db.query(Teacher)
        .filter(
            Teacher.user_id == current_user.id
        )
        .first()
    )

    if not teacher:
        raise HTTPException(
            status_code=404,
            detail="Teacher not found.",
        )

    assignment = (
        db.query(DescriptiveAssignment)
        .filter(
            DescriptiveAssignment.id == assignment_id,
            DescriptiveAssignment.teacher_id == teacher.id,
        )
        .first()
    )

    if not assignment:
        raise HTTPException(
            status_code=404,
            detail="Assignment not found.",
        )

    return {
        "id": assignment.id,
        "title": assignment.title,
        "subject_id": assignment.subject_id,
        "instructions": assignment.instructions,
        "status": assignment.status,
        "due_date": assignment.due_date,
        "created_at": assignment.created_at,
        "updated_at": assignment.updated_at,

        # -------------------------------------
        # Assigned Sections
        # -------------------------------------

        "sections": [
            {
                "id": mapping.section.id,

                "name": getattr(
                    mapping.section,
                    "name",
                    f"Section {mapping.section.id}",
                ),

                "section_name": getattr(
                    mapping.section,
                    "section_name",
                    None,
                ),

                "department": getattr(
                    mapping.section,
                    "department",
                    None,
                ),
            }

            for mapping in assignment.sections
        ],

        # -------------------------------------
        # Questions
        # -------------------------------------

        "questions": [
            {
                "id": question.id,
                "question_text": question.question_text,
                "max_marks": question.max_marks,
                "expected_answer": question.expected_answer,
                "evaluation_rubric": question.evaluation_rubric,
                "question_order": question.question_order,
            }

            for question in assignment.questions
        ],
    }


# =========================================================
# ASSIGN DESCRIPTIVE ASSIGNMENT TO SECTIONS
#
# Example:
#
# section_ids = [1, 2, 3]
#
# This:
# 1. Validates teacher owns sections
# 2. Creates assignment-section mappings
# 3. Changes assignment Draft -> Published
# 4. Finds students in those sections
# 5. Creates submissions for students
# 6. Sends assignment email to students
# =========================================================
@router.post(
    "/descriptive-assignments/{assignment_id}/assign"
)
def assign_descriptive_assignment(
    assignment_id: int,
    data: DescriptiveAssignmentSectionAssign,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    # =====================================================
    # FIND TEACHER
    # =====================================================

    teacher = (
        db.query(Teacher)
        .filter(
            Teacher.user_id == current_user.id
        )
        .first()
    )

    if not teacher:
        raise HTTPException(
            status_code=404,
            detail="Teacher not found.",
        )

    # =====================================================
    # FIND ASSIGNMENT
    # =====================================================

    assignment = (
        db.query(DescriptiveAssignment)
        .filter(
            DescriptiveAssignment.id == assignment_id,
            DescriptiveAssignment.teacher_id == teacher.id,
        )
        .first()
    )

    if not assignment:
        raise HTTPException(
            status_code=404,
            detail="Assignment not found.",
        )

    # =====================================================
    # VALIDATE SECTION SELECTION
    # =====================================================

    if not data.section_ids:
        raise HTTPException(
            status_code=400,
            detail="Select at least one section.",
        )

    # =====================================================
    # REMOVE DUPLICATE SECTION IDS
    # =====================================================

    section_ids = list(set(data.section_ids))

    # =====================================================
    # VERIFY TEACHER OWNS SELECTED SECTIONS
    # =====================================================

    teacher_sections = (
        db.query(TeacherSection)
        .filter(
            TeacherSection.teacher_id == teacher.id,
            TeacherSection.section_id.in_(section_ids),
        )
        .all()
    )

    allowed_section_ids = {
        item.section_id
        for item in teacher_sections
    }

    invalid_sections = [
        section_id
        for section_id in section_ids
        if section_id not in allowed_section_ids
    ]

    if invalid_sections:
        raise HTTPException(
            status_code=403,
            detail=(
                "You are not assigned to one or more "
                "selected sections."
            ),
        )

    # =====================================================
    # REMOVE OLD SECTION MAPPINGS
    #
    # This allows teacher to re-assign the assignment.
    # =====================================================

    db.query(
        DescriptiveAssignmentSection
    ).filter(
        DescriptiveAssignmentSection.assignment_id
        == assignment.id
    ).delete(
        synchronize_session=False
    )

    # =====================================================
    # ADD NEW SECTION MAPPINGS
    # =====================================================

    for section_id in section_ids:
        db.add(
            DescriptiveAssignmentSection(
                assignment_id=assignment.id,
                section_id=section_id,
            )
        )

    # =====================================================
    # PUBLISH ASSIGNMENT
    # =====================================================

    assignment.status = "Published"

    db.flush()

    # =====================================================
    # FIND STUDENTS FROM SELECTED SECTIONS
    # =====================================================

    students = (
        db.query(Student)
        .filter(
            Student.section_id.in_(section_ids),
            Student.is_active == True,
        )
        .all()
    )

    # =====================================================
    # CREATE / UPDATE SUBMISSIONS
    # =====================================================

    created_submissions = 0
    tokens_created = 0

    for student in students:

        submission = (
            db.query(DescriptiveSubmission)
            .filter(
                DescriptiveSubmission.assignment_id
                == assignment.id,

                DescriptiveSubmission.student_id
                == student.id,
            )
            .first()
        )

        # -------------------------------------------------
        # CREATE NEW SUBMISSION
        # -------------------------------------------------

        if not submission:

            submission = DescriptiveSubmission(
                assignment_id=assignment.id,
                student_id=student.id,
                status="In Progress",
                evaluation_status="Pending",
                token=secrets.token_urlsafe(32),
            )

            db.add(submission)

            created_submissions += 1
            tokens_created += 1

        # -------------------------------------------------
        # EXISTING SUBMISSION WITHOUT TOKEN
        # -------------------------------------------------

        elif not submission.token:

            submission.token = secrets.token_urlsafe(32)

            tokens_created += 1

    # =====================================================
    # SAVE SUBMISSIONS FIRST
    # =====================================================

    db.commit()

    # =====================================================
    # FRONTEND URL
    # =====================================================

    frontend_url = "http://localhost:5173"

    # =====================================================
    # SEND EMAILS
    # =====================================================

    emails_sent = 0
    emails_failed = 0

    for student in students:

        # -------------------------------------------------
        # VALIDATE STUDENT USER / EMAIL
        # -------------------------------------------------

        if not student.user or not student.user.email:

            emails_failed += 1
            continue

        # -------------------------------------------------
        # FIND STUDENT SUBMISSION
        # -------------------------------------------------

        submission = (
            db.query(DescriptiveSubmission)
            .filter(
                DescriptiveSubmission.assignment_id
                == assignment.id,

                DescriptiveSubmission.student_id
                == student.id,
            )
            .first()
        )

        if not submission or not submission.token:

            emails_failed += 1
            continue

        # -------------------------------------------------
# CREATE ASSIGNMENT LINK BASED ON TYPE
# -------------------------------------------------

        if assignment.assignment_type == "PDF":
           assignment_link = (
           f"{frontend_url}"
           f"/student/descriptive-pdf/"
           f"{submission.token}"
        )
        else:
         assignment_link = (
         f"{frontend_url}"
         f"/student/descriptive-assignments/token/"
         f"{submission.token}"
         )

        # -------------------------------------------------
        # SEND EMAIL
        # -------------------------------------------------

        try:

            send_quiz_email.invoke(
                {
                    "receiver_email": student.user.email,

                    "subject": (
                        f"New Assignment: "
                        f"{assignment.title}"
                    ),

                    "body": f"""
Hello {student.user.name},

A new descriptive assignment has been assigned to you.

----------------------------------------
Assignment Details
----------------------------------------

Assignment:
{assignment.title}

Instructions:
{assignment.instructions or "Please complete all questions carefully."}

Due Date:
{assignment.due_date or "No due date"}

----------------------------------------
Open Assignment
----------------------------------------

Click the link below to open and complete your assignment:

{assignment_link}

Please submit your answers before the due date.

Regards,
AI Training System
""",
                }
            )

            emails_sent += 1

        except Exception as e:

            print(
                f"Failed to send assignment email "
                f"to {student.user.email}: {e}"
            )

            emails_failed += 1

    # =====================================================
    # FINAL RESPONSE
    # =====================================================

    return {
        "assignment_id": assignment.id,
        "title": assignment.title,
        "status": assignment.status,
        "assigned_sections": section_ids,
        "total_students": len(students),
        "created_submissions": created_submissions,
        "tokens_created": tokens_created,
        "emails_sent": emails_sent,
        "emails_failed": emails_failed,
        "message": (
            "Descriptive assignment assigned successfully."
        ),
    }

@router.get(
    "/descriptive-assignments"
)
def get_teacher_descriptive_assignments(
    section_id: int | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        get_current_user
    ),
):
    teacher = (
        db.query(Teacher)
        .filter(
            Teacher.user_id == current_user.id
        )
        .first()
    )

    if not teacher:
        raise HTTPException(
            status_code=404,
            detail="Teacher not found.",
        )

    # -----------------------------------------
    # Base Query
    # -----------------------------------------

    query = (
        db.query(DescriptiveAssignment)
        .filter(
            DescriptiveAssignment.teacher_id
            == teacher.id
        )
    )

    # -----------------------------------------
    # Section Filter
    # -----------------------------------------

    if section_id is not None:

        query = (
            query
            .join(
                DescriptiveAssignmentSection,
                DescriptiveAssignmentSection.assignment_id
                == DescriptiveAssignment.id,
            )
            .filter(
                DescriptiveAssignmentSection.section_id
                == section_id
            )
        )

    # -----------------------------------------
    # Get assignments
    # -----------------------------------------

    assignments = (
        query
        .order_by(
            DescriptiveAssignment.created_at.desc()
        )
        .all()
    )

    result = []

    # -----------------------------------------
    # Build response
    # -----------------------------------------

    for assignment in assignments:

        sections = []

        for mapping in assignment.sections:

            section = mapping.section

            sections.append(
                {
                    "id": section.id,

                    "name": getattr(
                        section,
                        "name",
                        f"Section {section.id}",
                    ),

                    "section_name": getattr(
                        section,
                        "section_name",
                        None,
                    ),

                    "department": getattr(
                        section,
                        "department",
                        None,
                    ),
                }
            )

        result.append(
            {
                "id": assignment.id,

                "title": assignment.title,

                "subject_id": assignment.subject_id,

                "instructions": assignment.instructions,

                "status": assignment.status,

                "due_date": assignment.due_date,

                "created_at": assignment.created_at,

                "updated_at": assignment.updated_at,

                "sections": sections,

                "question_count": len(
                    assignment.questions
                ),
            }
        )

    return result


# =========================================================
# GET ASSIGNED STUDENTS FOR DESCRIPTIVE ASSIGNMENT
#
# Supports:
#
# GET
# /teachers/descriptive-assignments/{assignment_id}/assigned-students
#
# Optional:
#
# ?section_id=3
# ?search=Arun
#
# Returns only students who actually received
# this descriptive assignment.
# =========================================================
@router.get(
    "/descriptive-assignments/{assignment_id}/assigned-students"
)
def get_descriptive_assignment_assigned_students(
    assignment_id: int,
    section_id: int | None = None,
    search: str | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        get_current_user
    ),
):
    # -----------------------------------------
    # Find teacher
    # -----------------------------------------

    teacher = (
        db.query(Teacher)
        .filter(
            Teacher.user_id == current_user.id
        )
        .first()
    )

    if not teacher:
        raise HTTPException(
            status_code=404,
            detail="Teacher not found.",
        )

    # -----------------------------------------
    # Find assignment
    # -----------------------------------------

    assignment = (
        db.query(DescriptiveAssignment)
        .filter(
            DescriptiveAssignment.id == assignment_id,
            DescriptiveAssignment.teacher_id == teacher.id,
        )
        .first()
    )

    if not assignment:
        raise HTTPException(
            status_code=404,
            detail="Assignment not found.",
        )

    # -----------------------------------------
    # Get students who have submissions
    # -----------------------------------------

    query = (
        db.query(Student)
        .join(
            DescriptiveSubmission,
            DescriptiveSubmission.student_id
            == Student.id,
        )
        .join(
            User,
            User.id == Student.user_id,
        )
        .filter(
            DescriptiveSubmission.assignment_id
            == assignment.id
        )
    )

    # -----------------------------------------
    # Section filter
    # -----------------------------------------

    if section_id is not None:

        query = query.filter(
            Student.section_id == section_id
        )

    # -----------------------------------------
    # Search
    #
    # Searches:
    # - student name
    # - email
    # - roll number
    # -----------------------------------------

    if search:

        search_value = f"%{search.strip()}%"

        query = query.filter(
            (
                User.name.ilike(search_value)
            )
            |
            (
                User.email.ilike(search_value)
            )
            |
            (
                Student.roll_no.ilike(search_value)
            )
        )

    # -----------------------------------------
    # Remove duplicates
    # -----------------------------------------


    students = (
    query
    .group_by(Student.id, User.id)
    .order_by(User.name.asc())
    .all()
    )



    # -----------------------------------------
    # Build response
    # -----------------------------------------

    result = []

    for student in students:

        section = (
            db.query(Section)
            .filter(
                Section.id == student.section_id
            )
            .first()
        )

        section_name = (
            getattr(
                section,
                "name",
                None,
            )
            if section
            else f"Section {student.section_id}"
        )

        if (
            section
            and getattr(
                section,
                "section_name",
                None,
            )
        ):
            section_name = section.section_name

        result.append(
            {
                "student_id": student.id,

                "roll_no": student.roll_no,

                "student_name": student.user.name,

                "email": student.user.email,

                "section_id": student.section_id,

                "section_name": section_name,
            }
        )

    return {
        "assignment_id": assignment.id,
        "assignment_title": assignment.title,
        "total_students": len(result),
        "students": result,
    }

# =========================================================
# DESCRIPTIVE ASSIGNMENT SUBMISSIONS
# =========================================================

@router.get(
    "/descriptive-assignments/{assignment_id}/submissions"
)
def get_descriptive_assignment_submissions_api(
    assignment_id: int,
    current_user: User = Depends(
        get_current_user
    ),
    db: Session = Depends(get_db),
):
    return get_teacher_descriptive_assignment_submissions(
        db=db,
        current_user=current_user,
        assignment_id=assignment_id,
    )

# ============================================================
# DESCRIPTIVE ANALYTICS
# ============================================================

@router.get("/descriptive-analytics")
def get_descriptive_analytics(
    section_id: int | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    # ---------------------------------------------------------
    # VERIFY TEACHER
    # ---------------------------------------------------------
    teacher = (
        db.query(Teacher)
        .filter(
            Teacher.user_id == current_user.id
        )
        .first()
    )

    if not teacher:
        raise HTTPException(
            status_code=404,
            detail="Teacher not found.",
        )

    # ---------------------------------------------------------
    # TEACHER'S DESCRIPTIVE ASSIGNMENTS
    # ---------------------------------------------------------
    assignment_query = (
        db.query(DescriptiveAssignment)
        .filter(
            DescriptiveAssignment.teacher_id
            == teacher.id
        )
    )

    # If a section is selected, keep only assignments
    # that were assigned to that section.
    if section_id is not None:
        assignment_query = (
            assignment_query
            .join(
                DescriptiveAssignmentSection,
                DescriptiveAssignmentSection.assignment_id
                == DescriptiveAssignment.id,
            )
            .filter(
                DescriptiveAssignmentSection.section_id
                == section_id
            )
        )

    assignments = (
        assignment_query
        .distinct()
        .all()
    )

    assignment_ids = [
        assignment.id
        for assignment in assignments
    ]

    if not assignment_ids:
        return {
            "section_id": section_id,
            "total_assignments": 0,
            "total_students": 0,
            "evaluated_submissions": 0,
            "average_percentage": 0,
            "students": [],
        }

    # ---------------------------------------------------------
    # SUBMISSIONS
    # ---------------------------------------------------------
    submission_query = (
        db.query(DescriptiveSubmission)
        .join(
            Student,
            Student.id
            == DescriptiveSubmission.student_id,
        )
        .filter(
            DescriptiveSubmission.assignment_id.in_(
                assignment_ids
            )
        )
    )

    if section_id is not None:
        submission_query = submission_query.filter(
            Student.section_id == section_id
        )

    submissions = (
        submission_query
        .order_by(
            DescriptiveSubmission.id.asc()
        )
        .all()
    )

    # ---------------------------------------------------------
    # BUILD STUDENT PERFORMANCE
    #
    # A student's performance is the average percentage
    # across their evaluated descriptive submissions.
    # ---------------------------------------------------------
    student_map = {}

    for submission in submissions:
        student = submission.student

        if not student:
            continue

        if student.id not in student_map:
            student_map[student.id] = {
                "student_id": student.id,
                "student_name": (
                    student.user.name
                    if student.user
                    else "Unknown"
                ),
                "email": (
                    student.user.email
                    if student.user
                    else ""
                ),
                "roll_no": student.roll_no,
                "section_id": student.section_id,
                "section_name": None,
                "evaluated_submissions": 0,
                "pending_submissions": 0,
                "percentage_sum": 0.0,
            }

        item = student_map[student.id]

        status = str(
            submission.evaluation_status or ""
        ).lower()

        if status == "completed":
            percentage = float(
                submission.percentage or 0
            )

            item["percentage_sum"] += percentage
            item["evaluated_submissions"] += 1
        else:
            item["pending_submissions"] += 1

    # ---------------------------------------------------------
    # SECTION NAMES
    # ---------------------------------------------------------
    for item in student_map.values():
        section = (
            db.query(Section)
            .filter(
                Section.id
                == item["section_id"]
            )
            .first()
        )

        if section:
            item["section_name"] = (
                getattr(
                    section,
                    "section_name",
                    None,
                )
                or getattr(
                    section,
                    "name",
                    None,
                )
                or f"Section {section.id}"
            )
        else:
            item["section_name"] = (
                f"Section {item['section_id']}"
                if item["section_id"] is not None
                else "—"
            )

    # ---------------------------------------------------------
    # FINAL STUDENT LIST
    # ---------------------------------------------------------
    students = []

    for item in student_map.values():
        evaluated_count = item[
            "evaluated_submissions"
        ]

        average_percentage = (
            round(
                item["percentage_sum"]
                / evaluated_count,
                2,
            )
            if evaluated_count > 0
            else 0
        )

        students.append(
            {
                "student_id": item["student_id"],
                "student_name": item["student_name"],
                "email": item["email"],
                "roll_no": item["roll_no"],
                "section_id": item["section_id"],
                "section_name": item["section_name"],
                "evaluated_submissions": evaluated_count,
                "pending_submissions": item[
                    "pending_submissions"
                ],
                "average_percentage":
                    average_percentage,
            }
        )

    # Evaluated students first, then pending students.
    students.sort(
        key=lambda student: (
            student["evaluated_submissions"] > 0,
            student["average_percentage"],
        ),
        reverse=True,
    )

    evaluated_students = [
        student
        for student in students
        if student["evaluated_submissions"] > 0
    ]

    evaluated_submissions = sum(
        student["evaluated_submissions"]
        for student in students
    )

    average_percentage = (
        round(
            sum(
                student["average_percentage"]
                for student in evaluated_students
            )
            / len(evaluated_students),
            2,
        )
        if evaluated_students
        else 0
    )

    return {
        "section_id": section_id,
        "total_assignments": len(
            assignment_ids
        ),
        "total_students": len(students),
        "evaluated_submissions":
            evaluated_submissions,
        "average_percentage":
            average_percentage,
        "students": students,
    }

# =========================================================
# SINGLE DESCRIPTIVE SUBMISSION
# =========================================================

@router.get(
    "/descriptive-submissions/{submission_id}"
)
def get_descriptive_submission_api(
    submission_id: int,
    current_user: User = Depends(
        get_current_user
    ),
    db: Session = Depends(get_db),
):
    return get_teacher_descriptive_submission(
        db=db,
        current_user=current_user,
        submission_id=submission_id,
    )

@router.get(
    "/analytics/student-performance"
)
def teacher_student_performance(
    section_id: int | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
):
    teacher = (
        db.query(Teacher)
        .filter(
            Teacher.user_id == current_user.id
        )
        .first()
    )

    if not teacher:
        raise HTTPException(
            status_code=404,
            detail="Teacher profile not found.",
        )

    return get_student_performance(
        db,
        teacher.id,
        section_id,
    )


# =========================================================
# REVIEW DESCRIPTIVE SUBMISSION
# =========================================================

@router.put(
    "/descriptive-submissions/{submission_id}/review"
)
def review_descriptive_submission_api(
    submission_id: int,
    evaluation_data: DescriptiveSubmissionEvaluationUpdate,
    current_user: User = Depends(
        get_current_user
    ),
    db: Session = Depends(get_db),
):
    return review_teacher_descriptive_submission(
        db=db,
        current_user=current_user,
        submission_id=submission_id,
        evaluation_data=evaluation_data,
    )

@router.post("/descriptive-assignments/upload-pdf")
async def upload_descriptive_question_pdf(
    file: UploadFile = File(...),
    current_user: User = Depends(
        require_role(UserRole.TEACHER)
    ),
):
    # -----------------------------------------
    # Validate file
    # -----------------------------------------

    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="No file selected.",
        )

    if not file.filename.lower().endswith(".pdf"):
        raise HTTPException(
            status_code=400,
            detail="Only PDF files are allowed.",
        )

    # -----------------------------------------
    # Read file
    # -----------------------------------------

    file_content = await file.read()

    if not file_content:
        raise HTTPException(
            status_code=400,
            detail="Uploaded PDF is empty.",
        )

    # -----------------------------------------
    # Create upload directory
    # -----------------------------------------

    upload_dir = os.path.join(
        "uploads",
        "descriptive_questions",
    )

    os.makedirs(
        upload_dir,
        exist_ok=True,
    )

    # -----------------------------------------
    # Generate unique filename
    # -----------------------------------------

    unique_filename = (
        f"{uuid.uuid4().hex}_"
        f"{file.filename}"
    )

    file_path = os.path.join(
        upload_dir,
        unique_filename,
    )

    # -----------------------------------------
    # Save PDF
    # -----------------------------------------

    try:
        with open(file_path, "wb") as buffer:
            buffer.write(file_content)

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to save PDF: {str(e)}",
        )

    # -----------------------------------------
    # Return file information
    # -----------------------------------------

    file_url = (
        f"/uploads/descriptive_questions/"
        f"{unique_filename}"
    )

    return {
        "message": "Question PDF uploaded successfully.",
        "file_name": file.filename,
        "file_url": file_url,
    }

# ============================================================
# DOWNLOAD DESCRIPTIVE ASSIGNMENT SUBMISSIONS AS EXCEL
# ============================================================

@router.get(
    "/descriptive-assignments/{assignment_id}/submissions/download"
)
def download_descriptive_assignment_submissions(
    assignment_id: int,
    current_user: User = Depends(
        get_current_user
    ),
    db: Session = Depends(get_db),
):

    # --------------------------------------------------------
    # Verify teacher
    # --------------------------------------------------------

    teacher = (
        db.query(Teacher)
        .filter(
            Teacher.user_id == current_user.id
        )
        .first()
    )

    if not teacher:
        raise HTTPException(
            status_code=404,
            detail="Teacher not found."
        )

    # --------------------------------------------------------
    # Verify assignment ownership
    # --------------------------------------------------------

    assignment = (
        db.query(DescriptiveAssignment)
        .filter(
            DescriptiveAssignment.id == assignment_id,
            DescriptiveAssignment.teacher_id == teacher.id,
        )
        .first()
    )

    if not assignment:
        raise HTTPException(
            status_code=404,
            detail="Descriptive assignment not found."
        )

    # --------------------------------------------------------
    # Fetch submissions
    # --------------------------------------------------------

    submissions = (
        db.query(DescriptiveSubmission)
        .filter(
            DescriptiveSubmission.assignment_id
            == assignment.id
        )
        .order_by(
            DescriptiveSubmission.id.asc()
        )
        .all()
    )

    # --------------------------------------------------------
    # Create Excel workbook
    # --------------------------------------------------------

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Submissions"

    # --------------------------------------------------------
    # Headers
    # --------------------------------------------------------

    headers = [
        "Student Name",
        "Email",
        "Roll No",
        "Status",
        "Evaluation",
        "Obtained Marks",
        "Total Marks",
        "Percentage",
        "Submitted At",
        "Evaluated At",
    ]

    worksheet.append(headers)

    # --------------------------------------------------------
    # Data
    # --------------------------------------------------------

    for submission in submissions:

        student = submission.student

        student_name = (
            student.user.name
            if student
            and student.user
            else "Unknown"
        )

        student_email = (
            student.user.email
            if student
            and student.user
            else ""
        )

        roll_no = (
            student.roll_no
            if student
            else ""
        )

        worksheet.append([
            student_name,
            student_email,
            roll_no,
            submission.status or "",
            submission.evaluation_status or "",
            submission.obtained_marks
            if submission.obtained_marks is not None
            else 0,
            submission.total_marks
            if submission.total_marks is not None
            else 0,
            submission.percentage
            if submission.percentage is not None
            else 0,
            (
                submission.submitted_at.isoformat()
                if submission.submitted_at
                else ""
            ),
            (
                submission.evaluated_at.isoformat()
                if submission.evaluated_at
                else ""
            ),
        ])

    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    column_widths = {
        "A": 25,
        "B": 32,
        "C": 15,
        "D": 18,
        "E": 18,
        "F": 18,
        "G": 15,
        "H": 15,
        "I": 25,
        "J": 25,
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    # --------------------------------------------------------
    # Freeze header
    # --------------------------------------------------------

    worksheet.freeze_panes = "A2"

    # --------------------------------------------------------
    # Save workbook in memory
    # --------------------------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    # --------------------------------------------------------
    # Safe filename
    # --------------------------------------------------------

    safe_title = "".join(
        character
        if character.isalnum()
        else "_"
        for character in assignment.title
    )

    filename = (
        f"{safe_title}_Submissions.xlsx"
    )

    # --------------------------------------------------------
    # Return file
    # --------------------------------------------------------

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                f'attachment; filename="{filename}"'
        },
    )